#!/usr/bin/env python3
"""
Lleva a Google Drive las evidencias de la Bitácora de residentes.

1. Pide a la app lo pendiente (GET /api/sync): carpetas, archivos subidos y hojas de Excel.
2. Arma ese mismo árbol en una carpeta temporal (descarga archivos, genera los .xlsx).
3. Lo copia a Drive con UNA llamada a rclone.
4. Solo si rclone terminó bien, confirma a la app (POST /api/sync), que borra las copias temporales.
   Si algo falla no confirma: la próxima corrida lo reintenta.

Variables de entorno:
  APP_URL                 URL de la app, p. ej. https://bitacora-residentes.vercel.app
  SYNC_TOKEN              token compartido con la app
  BLOB_READ_WRITE_TOKEN   opcional: respaldo para descargar de Vercel Blob
  RCLONE_REMOTO           destino de rclone (por defecto "gdrive:")
  RCLONE_BIN              ejecutable de rclone (por defecto "rclone")
"""
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from urllib.parse import urlparse

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LOTES_MAXIMOS = 20  # la app entrega hasta 200 archivos por lote


class App:
    """Cliente HTTP de la app."""

    def __init__(self, url, token, token_blob=None):
        self.url = url.rstrip("/")
        self.http = requests.Session()
        self.http.headers["Authorization"] = f"Bearer {token}"
        self.token_blob = token_blob

    def pendientes_sync(self):
        r = self.http.get(f"{self.url}/api/sync", timeout=60)
        _comprobar(r)
        return r.json()

    def confirmar(self, confirmacion):
        r = self.http.post(f"{self.url}/api/sync", json=confirmacion, timeout=120)
        _comprobar(r)

    def descargar(self, url, destino):
        # Vercel Blob: la app entrega enlaces firmados temporales (no requieren token).
        # BLOB_READ_WRITE_TOKEN es opcional, solo como respaldo. En local se usa el token de la app.
        host = urlparse(url).hostname or ""
        if host.endswith(".blob.vercel-storage.com"):
            cabeceras = {"Authorization": f"Bearer {self.token_blob}"} if self.token_blob else {}
        else:
            cabeceras = {"Authorization": self.http.headers["Authorization"]}
        with requests.get(url, headers=cabeceras, stream=True, timeout=300) as r:
            _comprobar(r)
            with open(destino, "wb") as f:
                for trozo in r.iter_content(chunk_size=1024 * 1024):
                    f.write(trozo)


def _comprobar(r):
    if r.status_code >= 400:
        raise RuntimeError(f"{r.request.method} {r.url} respondió {r.status_code}: {r.text[:300]}")


def ruta_segura(base, ruta):
    """Ruta dentro de la carpeta temporal; rechaza rutas absolutas o con '..'."""
    partes = Path(ruta).parts
    if not ruta or Path(ruta).is_absolute() or ".." in partes:
        raise ValueError(f"Ruta inválida: {ruta!r}")
    return Path(base).joinpath(*partes)


def crear_xlsx(filas, destino):
    """Excel con encabezados en negrita sobre fondo verde oscuro, fila 1 fija y anchos ajustados."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Bitácora" if "Bitácora" in Path(destino).name else "Asistencia"
    for fila in filas:
        hoja.append(fila)
    encabezado = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="0F3D3E")
    for celda in hoja[1]:
        celda.font = encabezado
        celda.fill = fondo
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    for columna in range(1, (hoja.max_column or 0) + 1):
        largo = max((len(str(f[columna - 1])) for f in filas if len(f) >= columna), default=10)
        hoja.column_dimensions[get_column_letter(columna)].width = min(max(largo + 2, 12), 50)
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
    hoja.freeze_panes = "A2"
    libro.save(destino)


def armar(pendientes, base, descargar):
    """Reproduce en `base` el árbol de Drive con todo lo pendiente."""
    for carpeta in pendientes.get("carpetas", []):
        ruta_segura(base, carpeta).mkdir(parents=True, exist_ok=True)
    for archivo in pendientes.get("archivos", []):
        destino = ruta_segura(base, archivo["ruta"])
        destino.parent.mkdir(parents=True, exist_ok=True)
        descargar(archivo["url"], destino)
        if destino.stat().st_size != archivo["bytes"]:
            raise RuntimeError(f"El tamaño de {archivo['ruta']} no coincide; se reintentará en la próxima corrida.")
    for hoja in pendientes.get("hojas", []):
        destino = ruta_segura(base, hoja["ruta"])
        destino.parent.mkdir(parents=True, exist_ok=True)
        crear_xlsx(hoja["filas"], destino)


def sincronizar(app, rclone, remoto):
    """Un lote: armar → rclone → confirmar. Devuelve cuánto se llevó a Drive."""
    pendientes = app.pendientes_sync()
    resumen = {k: len(pendientes.get(k, [])) for k in ("carpetas", "archivos", "hojas")}
    if not any(resumen.values()):
        app.confirmar(pendientes.get("confirmacion", {}))  # solo registra la hora de la sincronización
        return resumen
    with tempfile.TemporaryDirectory(prefix="bitacora-") as tmp:
        armar(pendientes, tmp, app.descargar)
        comando = [rclone, "copy", tmp, remoto, "--create-empty-src-dirs", "--transfers", "4", "--retries", "3", "--stats-one-line", "-v"]
        if subprocess.run(comando).returncode != 0:
            raise RuntimeError("rclone no pudo copiar a Google Drive; no se confirma nada y se reintentará.")
    app.confirmar(pendientes["confirmacion"])
    return resumen


def main():
    faltan = [v for v in ("APP_URL", "SYNC_TOKEN") if not os.environ.get(v)]
    if faltan:
        sys.exit(f"Faltan variables de entorno: {', '.join(faltan)}")
    app = App(os.environ["APP_URL"], os.environ["SYNC_TOKEN"], os.environ.get("BLOB_READ_WRITE_TOKEN"))
    rclone = os.environ.get("RCLONE_BIN", "rclone")
    remoto = os.environ.get("RCLONE_REMOTO", "gdrive:")
    for lote in range(1, LOTES_MAXIMOS + 1):
        resumen = sincronizar(app, rclone, remoto)
        print(f"Lote {lote}: {resumen['carpetas']} carpetas, {resumen['archivos']} archivos, {resumen['hojas']} hojas.", flush=True)
        if resumen["archivos"] < 200 and resumen["hojas"] < 200:
            break
    print("Sincronización terminada.")


if __name__ == "__main__":
    main()
