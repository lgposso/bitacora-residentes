"""Pruebas del sincronizador. Ejecutar: .venv/bin/pytest sync"""
import os
import stat
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
import sincronizar as S  # noqa: E402

PENDIENTES = {
    "carpetas": ["MGB/Proyecto1/Pediatría", "MGB/Proyecto1/Pediatría/Ana/Evidencia Sesión 1/Registro Fotográfico"],
    "archivos": [{"ruta": "MGB/Proyecto1/Pediatría/Ana/Plan de clase.pdf", "url": "https://x.test/a", "bytes": 3}],
    "hojas": [{
        "ruta": "MGB/Proyecto1/Pediatría/Ana/Bitácora.xlsx",
        "filas": [["# Sesión", "Fecha", "Tema desarrollado"], [1, "29/09/2026", "Sepsis"]],
    }],
    "confirmacion": {"programas": True, "residentes": [1], "sesiones": [1], "archivos": [7], "bitacoras": [], "asistencias": []},
}


def test_crear_xlsx(tmp_path):
    destino = tmp_path / "Bitácora.xlsx"
    S.crear_xlsx([["# Sesión", "Tema"], [1, "Sepsis"], [2, "Asma"]], destino)
    hoja = load_workbook(destino).active
    assert [[c.value for c in fila] for fila in hoja.iter_rows()] == [["# Sesión", "Tema"], [1, "Sepsis"], [2, "Asma"]]
    assert hoja["A1"].font.bold
    assert hoja.freeze_panes == "A2"
    assert hoja.title == "Bitácora"


def test_ruta_segura_rechaza_salirse_de_la_carpeta(tmp_path):
    assert S.ruta_segura(tmp_path, "MGB/a/b.txt") == tmp_path / "MGB/a/b.txt"
    for mala in ["../x", "/etc/passwd", "MGB/../../x", ""]:
        with pytest.raises(ValueError):
            S.ruta_segura(tmp_path, mala)


def test_armar_crea_el_arbol(tmp_path):
    descargas = []

    def descargar(url, destino):
        descargas.append(url)
        destino.write_bytes(b"pdf")

    S.armar(PENDIENTES, tmp_path, descargar)
    assert (tmp_path / "MGB/Proyecto1/Pediatría/Ana/Evidencia Sesión 1/Registro Fotográfico").is_dir()
    assert (tmp_path / "MGB/Proyecto1/Pediatría/Ana/Plan de clase.pdf").read_bytes() == b"pdf"
    assert (tmp_path / "MGB/Proyecto1/Pediatría/Ana/Bitácora.xlsx").exists()
    assert descargas == ["https://x.test/a"]


def test_armar_falla_si_el_tamano_no_coincide(tmp_path):
    with pytest.raises(RuntimeError, match="tamaño"):
        S.armar(PENDIENTES, tmp_path, lambda url, destino: destino.write_bytes(b"x"))


def _rclone_falso(tmp_path, codigo):
    script = tmp_path / "rclone"
    script.write_text(f"#!/bin/sh\necho \"$@\" > {tmp_path}/args.txt\nexit {codigo}\n")
    script.chmod(script.stat().st_mode | stat.S_IEXEC)
    return str(script)


class AppFalsa:
    def __init__(self, pendientes):
        self.pendientes = pendientes
        self.confirmaciones = []

    def pendientes_sync(self):
        p, self.pendientes = self.pendientes, {"carpetas": [], "archivos": [], "hojas": [], "confirmacion": {}}
        return p

    def confirmar(self, c):
        self.confirmaciones.append(c)

    def descargar(self, url, destino):
        destino.write_bytes(b"pdf")


def test_sincronizar_confirma_solo_si_rclone_termina_bien(tmp_path):
    app = AppFalsa(PENDIENTES)
    resumen = S.sincronizar(app, _rclone_falso(tmp_path, 0), "gdrive:")
    assert resumen == {"carpetas": 2, "archivos": 1, "hojas": 1}
    assert app.confirmaciones == [PENDIENTES["confirmacion"]]
    args = (tmp_path / "args.txt").read_text()
    assert args.startswith("copy ") and "gdrive:" in args and "--create-empty-src-dirs" in args


def test_sincronizar_no_confirma_si_rclone_falla(tmp_path):
    app = AppFalsa(PENDIENTES)
    with pytest.raises(RuntimeError, match="rclone"):
        S.sincronizar(app, _rclone_falso(tmp_path, 1), "gdrive:")
    assert app.confirmaciones == []


def test_sin_pendientes_no_llama_a_rclone_pero_registra_la_hora(tmp_path):
    app = AppFalsa({"carpetas": [], "archivos": [], "hojas": [], "confirmacion": {"programas": False}})
    assert S.sincronizar(app, "/no/existe/rclone", "gdrive:") == {"carpetas": 0, "archivos": 0, "hojas": 0}
    assert app.confirmaciones == [{"programas": False}]
